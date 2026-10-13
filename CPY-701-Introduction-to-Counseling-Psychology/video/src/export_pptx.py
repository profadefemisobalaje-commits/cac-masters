#!/usr/bin/env python3
"""
CPY 701 - Introduction to Counseling Psychology
Editable PowerPoint export of the rendered slide decks.

Produces dist/CPY701-WNN-slides.pptx (16:9) mirroring slides/wNN.json so the
lecturer can edit, rebrand or present live from the same content that the
narrated video uses.
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x0B, 0x1B, 0x3A)
GOLD = RGBColor(0xC9, 0x96, 0x2E)
INK = RGBColor(0x16, 0x23, 0x3F)
SLATE = RGBColor(0x5A, 0x67, 0x84)
SURF = RGBColor(0xF7, 0xF8, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
DIST = os.path.join(ROOT, 'dist')


def add_base(prs, s, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY if dark else SURF
    # footer
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"CPY 701 · Introduction to Counseling Psychology      Master's Degree in Counseling Psychology      {s.get('idx','')} / {s.get('total','')}"
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(0xC9, 0xD3, 0xE8) if dark else SLATE
    if not dark:
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Pt(6))
        bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
        bar2 = slide.shapes.add_shape(1, 0, Pt(6), prs.slide_width, Pt(3))
        bar2.fill.solid(); bar2.fill.fore_color.rgb = GOLD; bar2.line.fill.background()
        chip = slide.shapes.add_shape(5, Inches(0.6), Inches(0.42), Inches(1.05), Inches(0.36))
        chip.fill.solid(); chip.fill.fore_color.rgb = NAVY; chip.line.fill.background()
        chip.text_frame.paragraphs[0].text = 'CPY 701'
        chip.text_frame.paragraphs[0].font.size = Pt(12)
        chip.text_frame.paragraphs[0].font.bold = True
        chip.text_frame.paragraphs[0].font.color.rgb = WHITE
        wk = slide.shapes.add_textbox(Inches(10.0), Inches(0.42), Inches(2.7), Inches(0.36))
        wk.text_frame.paragraphs[0].text = f"Week {s.get('week','')} · {s.get('module','')}"
        wk.text_frame.paragraphs[0].font.size = Pt(12)
        wk.text_frame.paragraphs[0].font.color.rgb = SLATE
        wk.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    return slide


def hdr(slide, s, y=Inches(1.05)):
    if s.get('kicker'):
        tb = slide.shapes.add_textbox(Inches(0.6), y, Inches(11), Inches(0.35))
        tb.text_frame.paragraphs[0].text = s['kicker'].upper()
        tb.text_frame.paragraphs[0].font.size = Pt(13)
        tb.text_frame.paragraphs[0].font.bold = True
        tb.text_frame.paragraphs[0].font.color.rgb = GOLD
    tb = slide.shapes.add_textbox(Inches(0.6), y + Inches(0.32), Inches(12.1), Inches(0.75))
    tb.text_frame.word_wrap = True
    tb.text_frame.paragraphs[0].text = s.get('title', '')
    tb.text_frame.paragraphs[0].font.size = Pt(30)
    tb.text_frame.paragraphs[0].font.bold = True
    tb.text_frame.paragraphs[0].font.color.rgb = NAVY


def body_items(slide, items, top=Inches(2.2), size=Pt(17)):
    tb = slide.shapes.add_textbox(Inches(0.6), top, Inches(12.1), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        lvl = it.get('lvl', 0)
        p.level = min(lvl, 2)
        run = p.add_run()
        if it.get('lead'):
            r0 = p.add_run(); r0.text = it['lead'] + ' '
            r0.font.bold = True; r0.font.size = size; r0.font.color.rgb = NAVY
        run.text = it.get('text', '')
        run.font.size = size if lvl == 0 else size - Pt(2)
        run.font.color.rgb = INK
        p.space_after = Pt(10)


def export(week):
    doc = json.load(open(os.path.join(ROOT, 'slides', f'{week}.json'))) if os.path.exists(os.path.join(ROOT, 'slides', f'{week}.json')) \
        else json.load(open(os.path.join(HERE, '..', 'slides', f'{week}.json')))
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slides = doc['slides']
    for i, s in enumerate(slides):
        s = dict(s); s['idx'] = i + 1; s['total'] = len(slides)
        s['week'] = s.get('week', doc.get('week')); s['module'] = s.get('module', doc.get('module'))
        lay = s['layout']
        dark = lay in ('title', 'section', 'quote', 'summary')
        slide = add_base(prs, s, dark)
        if lay == 'title':
            tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12), Inches(0.4))
            tb.text_frame.paragraphs[0].text = s['kicker'].upper()
            tb.text_frame.paragraphs[0].font.size = Pt(14)
            tb.text_frame.paragraphs[0].font.bold = True
            tb.text_frame.paragraphs[0].font.color.rgb = GOLD
            tb = slide.shapes.add_textbox(Inches(0.6), Inches(2.1), Inches(12), Inches(1.2))
            tb.text_frame.word_wrap = True
            tb.text_frame.paragraphs[0].text = s['title']
            tb.text_frame.paragraphs[0].font.size = Pt(40)
            tb.text_frame.paragraphs[0].font.bold = True
            tb.text_frame.paragraphs[0].font.color.rgb = WHITE
            tb = slide.shapes.add_textbox(Inches(0.6), Inches(3.2), Inches(12), Inches(0.6))
            tb.text_frame.paragraphs[0].text = s.get('subtitle', '')
            tb.text_frame.paragraphs[0].font.size = Pt(20)
            tb.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xC9, 0xD3, 0xE8)
            tb = slide.shapes.add_textbox(Inches(0.6), Inches(4.1), Inches(12), Inches(1.6))
            tf = tb.text_frame; tf.word_wrap = True
            for j, m in enumerate(s.get('meta', [])):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = m; p.font.size = Pt(14); p.font.color.rgb = RGBColor(0xC9, 0xD3, 0xE8)
        elif lay == 'quote':
            tb = slide.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(11.5), Inches(2.2))
            tf = tb.text_frame; tf.word_wrap = True
            tf.paragraphs[0].text = '\u201C' + s['quote'] + '\u201D'
            tf.paragraphs[0].font.size = Pt(28); tf.paragraphs[0].font.color.rgb = WHITE
            tb = slide.shapes.add_textbox(Inches(0.9), Inches(4.6), Inches(11), Inches(0.5))
            tb.text_frame.paragraphs[0].text = s['by']
            tb.text_frame.paragraphs[0].font.size = Pt(15); tb.text_frame.paragraphs[0].font.bold = True
            tb.text_frame.paragraphs[0].font.color.rgb = GOLD
        elif lay == 'summary':
            hdr(slide, {**s, 'kicker': 'Week in review'})
            for c in slide.shapes:
                pass
            body_items(slide, [{'text': t} for t in s['items']], top=Inches(2.1), size=Pt(17))
            for p in slide.shapes[-1].text_frame.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = RGBColor(0xE8, 0xEC, 0xF6)
        elif lay == 'table':
            hdr(slide, s)
            rows, cols = len(s['rows']) + 1, len(s['cols'])
            gt = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(2.1), Inches(12.1), Inches(0.5 * rows))
            tbl = gt.table
            for j, c in enumerate(s['cols']):
                tbl.columns[j].width = Emu(int(Inches(12.1) * c['w']))
                cell = tbl.cell(0, j)
                cell.text = c['t']
                cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
                cell.text_frame.paragraphs[0].font.bold = True
                cell.text_frame.paragraphs[0].font.color.rgb = WHITE
                cell.text_frame.paragraphs[0].font.size = Pt(14)
            for i, row in enumerate(s['rows'], 1):
                for j, v in enumerate(row):
                    cell = tbl.cell(i, j)
                    cell.text = v
                    cell.text_frame.paragraphs[0].font.size = Pt(13)
                    cell.text_frame.paragraphs[0].font.bold = (j == 0)
        elif lay == 'cards':
            hdr(slide, s)
            per = s.get('perRow', 3)
            for k, cd in enumerate(s['cards']):
                r, c = divmod(k, per)
                x = Inches(0.6 + c * 4.1); y = Inches(2.15 + r * 2.3)
                box = slide.shapes.add_shape(1, x, y, Inches(3.9), Inches(2.1))
                box.fill.solid(); box.fill.fore_color.rgb = WHITE
                box.line.color.rgb = RGBColor(0xD8, 0xDE, 0xEB)
                tf = box.text_frame; tf.word_wrap = True
                tf.margin_left = Inches(0.2); tf.margin_right = Inches(0.2); tf.margin_top = Inches(0.15)
                p = tf.paragraphs[0]; p.text = cd['title']; p.font.bold = True; p.font.size = Pt(16); p.font.color.rgb = NAVY
                p2 = tf.add_paragraph(); p2.text = cd['text']; p2.font.size = Pt(12); p2.font.color.rgb = SLATE
        elif lay == 'twocol':
            hdr(slide, s)
            for k, col in enumerate([s['left'], s['right']]):
                x = Inches(0.6 + k * 6.25)
                box = slide.shapes.add_shape(1, x, Inches(2.1), Inches(6.0), Inches(4.4))
                box.fill.solid(); box.fill.fore_color.rgb = WHITE
                box.line.color.rgb = RGBColor(0xD8, 0xDE, 0xEB)
                tf = box.text_frame; tf.word_wrap = True
                tf.margin_left = Inches(0.25); tf.margin_top = Inches(0.2)
                p = tf.paragraphs[0]; p.text = col['title']; p.font.bold = True; p.font.size = Pt(18); p.font.color.rgb = NAVY
                for it in col['items']:
                    p2 = tf.add_paragraph(); p2.text = '•  ' + it; p2.font.size = Pt(14); p2.font.color.rgb = INK
                    p2.space_after = Pt(8)
        elif lay in ('continuum', 'triangle', 'flow'):
            hdr(slide, s)
            items = s.get('nodes') or s.get('points') or s.get('steps') or []
            tb = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(12.1), Inches(3.6))
            tf = tb.text_frame; tf.word_wrap = True
            first = True
            for nd in items:
                p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
                p.text = f"{nd.get('title','')} — {nd.get('text','')}"
                p.font.size = Pt(16); p.font.color.rgb = INK; p.space_after = Pt(12)
            if s.get('note') or s.get('lede'):
                tb = slide.shapes.add_textbox(Inches(0.6), Inches(6.0), Inches(12.1), Inches(0.8))
                tb.text_frame.word_wrap = True
                tb.text_frame.paragraphs[0].text = s.get('note') or s.get('lede')
                tb.text_frame.paragraphs[0].font.size = Pt(14); tb.text_frame.paragraphs[0].font.italic = True
                tb.text_frame.paragraphs[0].font.color.rgb = SLATE
        elif lay == 'task':
            hdr(slide, s)
            body_items(slide, [{'text': t} for t in s['items']])
            if s.get('next'):
                tb = slide.shapes.add_textbox(Inches(0.6), Inches(6.2), Inches(12.1), Inches(0.7))
                tb.text_frame.word_wrap = True
                tb.text_frame.paragraphs[0].text = s['next']
                tb.text_frame.paragraphs[0].font.size = Pt(13); tb.text_frame.paragraphs[0].font.bold = True
                tb.text_frame.paragraphs[0].font.color.rgb = NAVY
        else:  # bullets and anything else
            hdr(slide, s)
            body_items(slide, s.get('items', []))
            if s.get('callout'):
                tb = slide.shapes.add_textbox(Inches(0.6), Inches(6.1), Inches(12.1), Inches(0.7))
                tb.text_frame.word_wrap = True
                tb.text_frame.paragraphs[0].text = f"{s['callout']['label']} {s['callout']['text']}"
                tb.text_frame.paragraphs[0].font.size = Pt(14)
                tb.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
    out = os.path.join(DIST, f'CPY701-{week.upper()}-slides.pptx')
    prs.save(out)
    print('saved', out)


if __name__ == '__main__':
    export(sys.argv[1] if len(sys.argv) > 1 else 'w01')
