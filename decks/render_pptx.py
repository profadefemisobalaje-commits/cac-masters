"""Render the layout primitives to a native PowerPoint file.

Everything is built on the blank layout so that no theme furniture can drift
in: the deck is drawn, not styled.  Text is set line by line from the
pre-wrapped layout, which is what keeps PowerPoint from re-flowing (and
overflowing) the design.
"""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from . import layout as L
from . import theme


def _rgb(h: str) -> RGBColor:
    r, g, b = theme.hex_rgb(h)
    return RGBColor(r, g, b)


def _pptx_face(family: str, weight: str, italic: bool):
    """Return (typeface, bold, italic) for a design weight."""
    if family == "display":
        if weight == "bold":
            return "EB Garamond", True, italic
        if weight == "semibold":
            return "EB Garamond SemiBold", False, italic
        if weight == "medium":
            return "EB Garamond Medium", False, italic
        return "EB Garamond", False, italic
    if weight == "light":
        return "Source Sans 3 Light", False, italic
    if weight == "semibold":
        return "Source Sans 3 Semibold", False, italic
    if weight == "bold":
        return "Source Sans 3", True, italic
    return "Source Sans 3", False, italic


def _alpha_fill(shape, alpha: float):
    if alpha >= 0.999:
        return
    spPr = shape._element.spPr
    srgb = spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int(round(alpha * 100000))))


def _alpha_run(run, alpha: float):
    if alpha >= 0.999:
        return
    srgb = run.font._rPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int(round(alpha * 100000))))


def _no_line(shape):
    shape.line.fill.background()


def _shape(shapes, kind, x, y, w, h):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if kind == "round" else MSO_SHAPE.RECTANGLE
    s = shapes.add_shape(st, Pt(x), Pt(y), Pt(max(w, 0.01)), Pt(max(h, 0.01)))
    s.shadow.inherit = False
    return s


def _draw_rect(shapes, op: L.Rect):
    s = _shape(shapes, "round" if op.radius > 0 else "rect",
               op.x, op.y, op.w, op.h)
    if op.radius > 0:
        s.adjustments[0] = min(0.5, op.radius / max(1e-6, min(op.w, op.h)))
    if op.fill:
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(op.fill)
        _alpha_fill(s, op.alpha)
    else:
        s.fill.background()
    if op.line:
        s.line.color.rgb = _rgb(op.line)
        s.line.width = Pt(op.lw)
    else:
        _no_line(s)
    return s


def _draw_ellipse(shapes, op: L.Ellipse):
    s = shapes.add_shape(MSO_SHAPE.OVAL, Pt(op.x), Pt(op.y), Pt(max(op.w, 0.01)), Pt(max(op.h, 0.01)))
    s.shadow.inherit = False
    if op.fill:
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(op.fill)
        _alpha_fill(s, op.alpha)
    else:
        s.fill.background()
    if op.line:
        s.line.color.rgb = _rgb(op.line)
        s.line.width = Pt(op.lw)
    else:
        _no_line(s)
    return s


def _draw_rule(shapes, op: L.Rule):
    if abs(op.y2 - op.y1) < 0.01:                 # horizontal
        x, y = min(op.x1, op.x2), op.y1
        w, h = abs(op.x2 - op.x1), max(op.lw, 0.35)
    elif abs(op.x2 - op.x1) < 0.01:               # vertical
        x, y = op.x1, min(op.y1, op.y2)
        w, h = max(op.lw, 0.35), abs(op.y2 - op.y1)
    else:
        conn = shapes.add_connector(1, Pt(op.x1), Pt(op.y1), Pt(op.x2), Pt(op.y2))
        conn.line.color.rgb = _rgb(op.color)
        conn.line.width = Pt(op.lw)
        if op.dash:
            conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return conn
    s = _shape(shapes, "rect", x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(op.color)
    _no_line(s)
    if op.dash:
        s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return s


_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def _draw_text(shapes, op: L.Text):
    h = op.height + op.size * 0.6
    tb = shapes.add_textbox(Pt(op.x), Pt(op.y - op.size * 0.16),
                            Pt(max(op.w, 1)), Pt(max(h, 1)))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    face, bold, italic = _pptx_face(op.family, op.weight, op.italic)
    for i, line in enumerate(op.lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[op.align]
        p.line_spacing = Pt(op.line_height)
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = line
        f = r.font
        f.name = face
        f.size = Pt(op.size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = _rgb(op.color)
        if op.tracking:
            f._rPr.set("spc", str(int(round(op.tracking * 100))))
        _alpha_run(r, op.alpha)
    return tb


def _draw_image(shapes, op: L.Image):
    return shapes.add_picture(op.path, Pt(op.x), Pt(op.y), Pt(op.w), Pt(op.h))


def render(deck, path: Path, title: str | None = None,
           subtitle: str | None = None) -> Path:
    pal = deck.course.palette
    prs = Presentation()
    prs.slide_width = Pt(theme.W)
    prs.slide_height = Pt(theme.H)
    blank = prs.slide_layouts[6]

    for s in deck.slides:
        slide = prs.slides.add_slide(blank)

        # Background: solid base + optional texture plate.
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(s.bg)
        if s.texture and Path(s.texture).exists():
            pic = slide.shapes.add_picture(str(s.texture), 0, 0,
                                           Pt(theme.W), Pt(theme.H))
            # Send to back (immediately after the background).
            sp = pic._element
            sp.getparent().remove(sp)
            slide.shapes._spTree.insert(2, sp)

        # Decorative hairlines drawn first so text sits above.
        for op in s.ops:
            if op.kind == "rect":
                _draw_rect(slide.shapes, op)
            elif op.kind == "ellipse":
                _draw_ellipse(slide.shapes, op)
            elif op.kind == "rule":
                _draw_rule(slide.shapes, op)

        for op in s.ops:
            if op.kind == "image":
                _draw_image(slide.shapes, op)
            elif op.kind == "text":
                _draw_text(slide.shapes, op)

        if s.notes:
            slide.notes_slide.notes_text_frame.text = s.notes

    cp = prs.core_properties
    cp.title = title or f"{deck.course.code} — {deck.course.title}"
    cp.subject = subtitle or deck.course.programme
    cp.author = "Master of Marriage and Family Therapy"
    cp.category = "Lecture deck"
    cp.comments = ("Designed for the MFT lecture series. "
                   "Set in EB Garamond and Source Sans 3.")

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path
